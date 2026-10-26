# Converte PDFs do Itaú (extrato e fatura) e a planilha do C6 em um arquivo padronizado
# lido pela página "Contas da Família". Uso: python3 normalizar.py  (procura nas subpastas de ..)
import os,re,json,csv,hashlib,subprocess,datetime,collections,sys
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
from itau_fatura import parse as parse_fatura
import openpyxl
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)));OUT=os.path.dirname(os.path.abspath(__file__))
def num(s): return float(s.replace('.','').replace(',','.'))
def hid(*a): return hashlib.sha1("|".join(map(str,a)).encode()).hexdigest()[:10]
rows=[];files=[];faturas=[];saldos=collections.defaultdict(dict);avisos=[];seen_md5={}
def add(**k): rows.append(k)
def md5(p): return hashlib.md5(open(p,'rb').read()).hexdigest()
def itau_extrato(p,fn):
    t=subprocess.run(["pdftotext","-layout",p,"-"],capture_output=True,text=True).stdout
    per=re.search(r'período de visualização:\s*(\d\d)/(\d\d)/(\d{4})\s+até\s+(\d\d)/(\d\d)/(\d{4})',t)
    ini=f"{per.group(3)}-{per.group(2)}-{per.group(1)}";fim=f"{per.group(6)}-{per.group(5)}-{per.group(4)}"
    sal=re.search(r'saldo em conta[\s\S]*?R\$\s*(-?[\d.]+,\d\d)',t)
    occ=collections.Counter();n=0
    for l in t.splitlines():
        m=re.match(r'^(\d\d)/(\d\d)/(\d{4})\s+(.+?)\s{2,}(-?[\d.]+,\d\d)\s*$',l.rstrip())
        if not m: continue
        d=f"{m.group(3)}-{m.group(2)}-{m.group(1)}";desc=re.sub(r'\s+',' ',m.group(4)).strip();v=num(m.group(5))
        if desc.startswith('SALDO DO DIA'):
            mo=d[:7]
            if mo not in saldos['itau-cc'] or d>=saldos['itau-cc'][mo][1]: saldos['itau-cc'][mo]=[v,d]
            continue
        if d<ini or d>fim: continue
        k=(d,desc,v);occ[k]+=1
        tipo='encargo' if re.search(r'JUROS|IOF|MULTA',desc) else 'divida' if 'RECLASSIF' in desc else 'mov'
        add(id='I'+hid(d,desc,v,occ[k]),conta='itau-cc',titular='Silvano',data=d,valor=v,descricao=desc,categoria_banco='',portador='',cartao_final='',parcela='',vencimento='',tipo=tipo,arquivo=fn);n+=1
    # conciliação: saldo anterior ao período + lançamentos = último saldo do dia
    sd=[(f"{m.group(3)}-{m.group(2)}-{m.group(1)}",num(m.group(4))) for m in re.finditer(r'^(\d\d)/(\d\d)/(\d{4})\s+SALDO DO DIA\s+(-?[\d.]+,\d\d)',t,re.M)]
    dentro=[x for x in sd if ini<=x[0]<=fim];antes=[x for x in sd if x[0]<ini]
    ok=None
    if dentro and antes:
        mine=[r for r in rows if r['arquivo']==fn];ult=max(r['data'] for r in mine);soma=sum(r['valor'] for r in mine)
        fimv=max(x for x in dentro if x[0]<=ult);ok=abs(max(antes)[1]+soma-fimv[1])<0.02
    files.append(dict(arquivo=fn,conta='itau-cc',tipo='Extrato Itaú',periodo=[ini,fim],lancamentos=n,validado=ok,saldo_cabecalho=(num(sal.group(1)) if sal else None),saldo_cabecalho_data=re.search(r'emitido em:\s*(\d\d)/(\d\d)/(\d{4})',t).expand(r'\3-\2-\1')))
def itau_fatura(p,fn):
    r=parse_fatura(p)
    for i,t in enumerate(r['tx']):
        own=t['own'] or '';tit='Gisele' if own.upper().startswith('GISE') else 'Silvano'
        comp=r['fech'] if t['parc'] else t['date']
        add(id='F'+hid(r['venc'],t['date'],t['desc'],t['v'],t['parc'],t['fin'],i),conta='itau-cartao',titular=tit,data=comp,valor=-t['v'],descricao=t['desc'],categoria_banco=t['cat'],portador=own,cartao_final=t['fin'] or '',parcela=t['parc'],vencimento=r['venc'],tipo=('encargo' if t['sec']=='prod' or t['desc'].startswith('IOF') else 'mov'),arquivo=fn)
    if r['enc']>0:
        add(id='F'+hid(r['venc'],'encargos'),conta='itau-cartao',titular='Silvano',data=r['fech'],valor=-r['enc'],descricao='Encargos da fatura (juros do rotativo, mora, multa, IOF)',categoria_banco='',portador='',cartao_final='',parcela='',vencimento=r['venc'],tipo='encargo',arquivo=fn)
    ok=abs(r['soma']-r['L'])<0.02
    if not ok: avisos.append(f"{fn}: soma dos lançamentos {r['soma']:.2f} difere do total impresso {r['L']:.2f}")
    faturas.append(dict(conta='itau-cartao',vencimento=r['venc'],fechamento=r['fech'],total=r['total'],lancamentos=r['L'],encargos=r['enc'],anterior=r['ant'],pago=r['pago'],validado=ok,arquivo=fn))
    files.append(dict(arquivo=fn,conta='itau-cartao',tipo='Fatura Itaú',periodo=[r['venc'],r['venc']],lancamentos=len(r['tx']),validado=ok))
def c6_xlsx(p,fn):
    ws=openpyxl.load_workbook(p,data_only=True).active;occ=collections.Counter();n=0;ini=fim=None;last={}
    for r in ws.iter_rows(min_row=10,values_only=True):
        if not isinstance(r[0],datetime.datetime): continue
        d=r[0].strftime('%Y-%m-%d');tit=(r[2] or '').strip();det=(r[3] or '').strip();v=round((r[4] or 0)-(r[5] or 0),2)
        desc=tit if re.match(r'^(Pix|Recorr|Devol|Transf)',tit,re.I) else (det if tit.upper().startswith('DEBITO DE CARTAO') or tit.upper().startswith('EST DEBITO') else (tit+(' · '+det if det and det.upper()!=tit.upper() else '')))
        desc=re.sub(r'\s+',' ',desc).strip()
        tipo='pgto_fatura' if 'PGTO FAT' in tit.upper() else 'investimento' if re.search(r'CDB|TESOURO',tit,re.I) else 'mov'
        k=(d,desc,v);occ[k]+=1
        add(id='S'+hid(d,desc,v,occ[k]),conta='c6-cc',titular='Gisele',data=d,valor=v,descricao=desc,categoria_banco=tit,portador='',cartao_final='',parcela='',vencimento='',tipo=tipo,arquivo=fn);n+=1
        ini=min(ini or d,d);fim=max(fim or d,d)
        if r[6] is not None: saldos['c6-cc'][d[:7]]=[float(r[6]),d]
    # conciliação: saldo do dia (por data contábil) x soma dos lançamentos
    import itertools
    rs=[r for r in ws.iter_rows(min_row=10,values_only=True) if isinstance(r[0],datetime.datetime)]
    g0=rs[0];ini_saldo=None;difs=[]
    first=[r for r in rs if r[1]==g0[1]];ini_saldo=round(first[-1][6]-sum((r[4] or 0)-(r[5] or 0) for r in first),2)
    bal=ini_saldo
    for dc,g in itertools.groupby(rs,key=lambda r:r[1]):
        g=list(g);bal+=sum((r[4] or 0)-(r[5] or 0) for r in g)
        if g[-1][6] is not None and abs(bal-g[-1][6])>0.01:
            difs.append([dc.strftime('%Y-%m-%d'),round(g[-1][6]-bal,2)]);bal=g[-1][6]
    for d,v in difs: avisos.append(f"C6: em {d[8:]}/{d[5:7]} o saldo do banco difere em R$ {v:.2f} da soma dos lançamentos (movimento ausente no extrato exportado)")
    files.append(dict(arquivo=fn,conta='c6-cc',tipo='Extrato C6 (planilha)',periodo=[ini,fim],lancamentos=n,validado=(len(difs)==0),saldo_inicial=ini_saldo,diferencas=difs))
for dp,dn,fn in os.walk(ROOT):
    dn[:]=[x for x in dn if not x.startswith('_')]
    for f in sorted(fn):
        p=os.path.join(dp,f);low=f.lower()
        if not low.endswith(('.pdf','.xlsx','.xls')): continue
        h=md5(p)
        if h in seen_md5: avisos.append(f"{f} é cópia idêntica de {seen_md5[h]} e foi ignorado"); files.append(dict(arquivo=f,conta='',tipo='Duplicado',ignorado=True));continue
        seen_md5[h]=f
        try:
            if low.endswith('.xlsx') and 'c6' in low: c6_xlsx(p,f)
            elif low.endswith('.pdf'):
                head=subprocess.run(["pdftotext","-l","1",p,"-"],capture_output=True,text=True).stdout
                if 'extrato conta / lançamentos' in head or 'extrato conta' in head and 'Itaú' in head or re.search(r'itau_extrato',low): itau_extrato(p,f)
                elif 'Resumo da fatura' in head or 'Total desta fatura' in head: itau_fatura(p,f)
                elif 'C6' in head or 'c6' in low:
                    files.append(dict(arquivo=f,conta='c6-cc',tipo='PDF do C6 (mesmos dados da planilha)',ignorado=True))
                else: avisos.append(f"{f}: formato não reconhecido"); files.append(dict(arquivo=f,conta='',tipo='Não reconhecido',ignorado=True))
            else: avisos.append(f"{f}: planilha não reconhecida"); files.append(dict(arquivo=f,conta='',tipo='Não reconhecido',ignorado=True))
        except Exception as e:
            avisos.append(f"{f}: erro ao ler ({e})"); files.append(dict(arquivo=f,conta='',tipo='Erro',ignorado=True))
# lançamentos manuais (transcritos de capturas de tela etc.), em _app/manuais/*.csv
MAN=os.path.join(OUT,'manuais')
if os.path.isdir(MAN):
    for f in sorted(os.listdir(MAN)):
        if not f.lower().endswith('.csv'): continue
        L=open(os.path.join(MAN,f),encoding='utf-8').read().splitlines();meta_m={}
        for l in L:
            if l.startswith('#saldo_'): k,v,d=l[1:].split(';')[:3];meta_m[k]=(float(v),d)
        body=[l for l in L if l and not l.startswith('#')];hdr=body[0].split(';');occ=collections.Counter();soma=0;n=0;datas=[];contas=set()
        for l in body[1:]:
            c=dict(zip(hdr,l.split(';')));v=float(c['valor']);d=c['data'];desc=c['descricao'].strip();k=(c['conta'],d,desc,v);occ[k]+=1
            tipo='encargo' if re.search(r'JUROS|IOF|MULTA',desc) else 'mov'
            add(id='M'+hid(*k,occ[k]),conta=c['conta'],titular=c.get('titular','Silvano'),data=d,valor=v,descricao=desc,categoria_banco='',portador='',cartao_final='',parcela='',vencimento='',tipo=tipo,arquivo=f)
            soma+=v;n+=1;datas.append(d);contas.add(c['conta'])
        ok=None
        if 'saldo_inicial' in meta_m and 'saldo_final' in meta_m:
            ok=abs(meta_m['saldo_inicial'][0]+soma-meta_m['saldo_final'][0])<0.02
            if not ok: avisos.append(f"{f}: lançamentos manuais não fecham com os saldos informados")
        per=[(lambda d,n:(datetime.date.fromisoformat(d)+datetime.timedelta(days=n)).isoformat())(meta_m['saldo_inicial'][1],1) if 'saldo_inicial' in meta_m else min(datas), meta_m['saldo_final'][1] if 'saldo_final' in meta_m else max(datas)]
        files.append(dict(arquivo=f,conta=sorted(contas)[0] if contas else '',tipo='Lançamentos manuais (capturas de tela)',periodo=per,lancamentos=n,validado=ok))
rows.sort(key=lambda r:(r['data'],r['conta'],r['valor']))
cols=['id','conta','titular','data','valor','descricao','categoria_banco','portador','cartao_final','parcela','vencimento','tipo','arquivo']
with open(os.path.join(OUT,'lancamentos_extra.csv'),'w',newline='',encoding='utf-8') as fh:
    w=csv.writer(fh,delimiter=';');w.writerow(cols)
    for r in rows: w.writerow([('%.2f'%r[c]) if c=='valor' else str(r[c]).replace(';',',') for c in cols])
meta=dict(gerado_em=datetime.datetime.now().isoformat(timespec='seconds'),arquivos=files,faturas=sorted(faturas,key=lambda f:f['vencimento']),
  saldos={k:{m:v for m,v in sorted(d.items())} for k,d in saldos.items()},avisos=avisos,total_lancamentos=len(rows))
json.dump(meta,open(os.path.join(OUT,'contas_extra.json'),'w',encoding='utf-8'),ensure_ascii=False,indent=1)
print(len(rows),'lançamentos;',len([f for f in files if not f.get('ignorado')]),'arquivos processados');print('\n'.join(avisos))
